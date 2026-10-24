
import re
from pathlib import Path

import fitz  
from docx import Document


MAX_FILE_SIZE = 5 * 1024 * 1024

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}

SKILL_KEYWORDS = {
    "Python": ["python"],
    "Java": ["java"],
    "C++": ["c++"],
    "JavaScript": ["javascript", "js"],
    "SQL": ["sql", "mysql", "postgresql"],
    "Machine Learning": ["machine learning", "scikit-learn", "sklearn"],
    "Deep Learning": ["deep learning", "tensorflow", "pytorch", "keras"],
    "Data Analysis": ["data analysis", "pandas", "numpy"],
    "Natural Language Processing": [
        "natural language processing", "nlp", "transformers"
    ],
    "Computer Vision": ["computer vision", "opencv", "image classification"],
    "FastAPI": ["fastapi"],
    "Flask": ["flask"],
    "Django": ["django"],
    "HTML/CSS": ["html", "css"],
    "React": ["react", "react.js"],
    "Git": ["git", "github"],
    "Docker": ["docker"],
    "AWS": ["aws", "amazon web services"],
}


def extract_resume_text(
    file_bytes: bytes,
    filename: str,
) -> str:
    """Extract text from a PDF or DOCX resume."""

    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Only PDF and DOCX resumes are supported.")

    if not file_bytes:
        raise ValueError("The uploaded file is empty.")

    if len(file_bytes) > MAX_FILE_SIZE:
        raise ValueError("Resume must be 5 MB or smaller.")

    if extension == ".pdf":
        text_parts = []

        with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
            if pdf.page_count > 20:
                raise ValueError("Resume must not exceed 20 pages.")

            for page in pdf:
                text_parts.append(page.get_text())

        text = "\n".join(text_parts)

    else:
        import io

        document = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in document.paragraphs]

        for table in document.tables:
            for row in table.rows:
                paragraphs.append(
                    " ".join(cell.text for cell in row.cells)
                )

        text = "\n".join(paragraphs)

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()

    if len(text) < 30:
        raise ValueError(
            "Not enough text could be extracted. "
            "If this is a scanned PDF, OCR support is needed."
        )

    return text


def extract_skills(resume_text: str) -> list[str]:
    """Find skills mentioned in the resume using a starter vocabulary."""

    text = resume_text.lower()
    found_skills = []

    for skill, keywords in SKILL_KEYWORDS.items():
        for keyword in keywords:
            pattern = (
                r"(?<![a-z0-9+#.])"
                + re.escape(keyword)
                + r"(?![a-z0-9+#.])"
            )

            if re.search(pattern, text):
                found_skills.append(skill)
                break

    return found_skills


def analyse_resume(file_bytes: bytes, filename: str) -> dict:
    """Extract resume text, skills and a predicted job category."""

    text = extract_resume_text(file_bytes, filename)
    skills = extract_skills(text)

    predicted_role = None
    classifier_status = "not_configured"

    try:
        try:
            from models.classifier import predict_role
        except ImportError:
            from classifier import predict_role

        predicted_role = predict_role(text)
        classifier_status = "success"

    except (ImportError, FileNotFoundError):

        pass

    return {
        "filename": Path(filename).name,
        "extracted_text": text,
        "predicted_role": predicted_role,
        "skills": skills,
        "skill_count": len(skills),
        "classifier_status": classifier_status,
    }
