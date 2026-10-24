import os
import sys
import json
from pathlib import Path
from dotenv import load_dotenv

load_dotenv()

import math
import requests
import numpy as np
from datetime import date

import hashlib
import logging
from datetime import datetime, timezone
from starlette.concurrency import run_in_threadpool


from pathlib import Path
import pandas as pd
from authlib.integrations.starlette_client import OAuth
from fastapi import FastAPI, Request
from fastapi.responses import HTMLResponse, RedirectResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.middleware.sessions import SessionMiddleware
from pydantic import BaseModel

import re
import fitz

from docx import Document
from fastapi import UploadFile, File, HTTPException

import uuid
from fastapi import UploadFile, File, HTTPException

from models.salary_ai import apply_ai_salaries

from models.career_detail import build_detail

from typing import Literal
from pydantic import Field
from models.plan_ai import generate_plan

from models.sdg_ai import analyse as sdg_analyse

from fastapi.responses import Response

from models.lesson_ai import generate_lesson, generate_quiz


BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data"

RESUME_STORAGE_DIR = BASE_DIR / "data" / "saved_resumes"
RESUME_STORAGE_DIR.mkdir(parents=True, exist_ok=True)

QUIZ_RESULT_DIR = BASE_DIR / "data" / "quiz_results"
QUIZ_RESULT_DIR.mkdir(parents=True, exist_ok=True)

RESKILL_DIR = DATA_DIR / "reskilling"
RESKILL_DIR.mkdir(parents=True, exist_ok=True)


LESSON_DIR = DATA_DIR / "lessons"
LESSON_DIR.mkdir(parents=True, exist_ok=True)


SDG_DIR = DATA_DIR / "sdg"
SDG_DIR.mkdir(parents=True, exist_ok=True)

MAX_RESUME_SIZE = 5 * 1024 * 1024  
ALLOWED_RESUME_EXTENSIONS = {".pdf", ".docx"}

SC_DIR = DATA_DIR / "scenarios"
SC_DIR.mkdir(parents=True, exist_ok=True)
SC_EXPOSURE_FILE = SC_DIR / "exposure_cache.json"
SC_YEARS = 5
SC_SPEED = {"moderate": 1.0, "fast": 2.0} 
DEEPSEEK_MODEL = "deepseek-chat"

SETTINGS_DIR = DATA_DIR / "settings"
SETTINGS_DIR.mkdir(parents=True, exist_ok=True)

app = FastAPI()

if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from models.career_matcher import (
    QUESTIONS,
    QUESTION_MAP,
    RIASEC_COLUMNS,
    RIASEC_LETTERS,
    score_riasec,
    match_careers,
)


QUESTION_CATEGORY = {q: category for category, qs in QUESTION_MAP.items() for q in qs}


class QuizSubmission(BaseModel):
    answers: dict[int, int]


RESUME_MODEL_DIR = BASE_DIR / "models" / "distilbert_resume_classifier_v2"
RESUME_LABEL_ENCODER = (
    BASE_DIR / "models" / "distilbert_resume_classifier_v2" / "label_encoder.pkl"
)
MAX_RESUME_SIZE = 5 * 1024 * 1024
ALLOWED_EXTENSIONS = {".pdf", ".docx"}

_device = None
_resume_tokenizer = None
_resume_model = None
_resume_label_encoder = None


def load_resume_model():
    """Load the trained model once, when first needed."""
    global _resume_tokenizer, _resume_model, _resume_label_encoder, _device

    if _resume_model is not None:
        return

    if not RESUME_MODEL_DIR.exists():
        raise HTTPException(
            status_code=503,
            detail="Resume classification is not configured. Add the trained "
            "model to models/distilbert_resume_classifier_v2.",
        )

    if not RESUME_LABEL_ENCODER.exists():
        raise HTTPException(
            status_code=503,
            detail="Resume classification is not configured: the trained "
            "model's label_encoder.pkl is missing.",
        )

    try:
        import joblib
        import torch
        from transformers import AutoTokenizer, AutoModelForSequenceClassification
    except ImportError as exc:
        raise HTTPException(
            status_code=503,
            detail="Resume classification dependencies are missing. "
            "Install requirements-resume.txt after adding the trained model.",
        ) from exc

    _device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    _resume_tokenizer = AutoTokenizer.from_pretrained(str(RESUME_MODEL_DIR))
    _resume_model = AutoModelForSequenceClassification.from_pretrained(
        str(RESUME_MODEL_DIR)
    )
    _resume_label_encoder = joblib.load(RESUME_LABEL_ENCODER)

    _resume_model.to(_device)
    _resume_model.eval()


def extract_resume_text(filename: str, content: bytes) -> str:
    """Extract text locally from a PDF or DOCX file."""
    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        import pymupdf

        with pymupdf.open(stream=content, filetype="pdf") as pdf:
            text = "\n".join(page.get_text() for page in pdf)

    elif extension == ".docx":
        import io

        document = Document(io.BytesIO(content))
        paragraphs = [p.text for p in document.paragraphs]

        for table in document.tables:
            for row in table.rows:
                paragraphs.extend(cell.text for cell in row.cells)

        text = "\n".join(paragraphs)

    else:
        raise HTTPException(
            status_code=400, detail="Only PDF and DOCX files are supported."
        )

    text = re.sub(r"\s+", " ", text).strip()

    if len(text) < 50:
        raise HTTPException(
            status_code=422,
            detail=(
                "Not enough readable text was found. "
                "This may be a scanned PDF; OCR is not yet enabled "
                "for uploaded resumes."
            ),
        )

    return text


def _normalise_label(name: str) -> str:
    """'Sales resumes' and 'Sales' are the same category."""
    n = str(name).lower().replace("_", " ")
    n = re.sub(r"\bresumes?\b", "", n)
    return re.sub(r"[^a-z0-9]", "", n)


_LABEL_GROUPS = None


def _get_label_groups():
    """Group the model's class indices that mean the same category."""
    global _LABEL_GROUPS
    if _LABEL_GROUPS is None:
        groups = {}
        for idx, name in enumerate(_resume_label_encoder.classes_):
            groups.setdefault(_normalise_label(name), []).append(idx)
        _LABEL_GROUPS = [
            (min((str(_resume_label_encoder.classes_[i]) for i in idxs), key=len), idxs)
            for idxs in groups.values()
        ]
    return _LABEL_GROUPS


def predict_resume_category(text: str):
    """Classify resume using RunPod Serverless."""
    api_key = os.getenv("RUNPOD_API_KEY")
    endpoint_id = os.getenv("RUNPOD_ENDPOINT_ID")

    if not api_key or not endpoint_id:
        raise HTTPException(
            status_code=503,
            detail="RunPod API credentials are not configured.",
        )

    url = f"https://api.runpod.ai/v2/{endpoint_id}/runsync"

    try:
        response = requests.post(
            url,
            headers={
                "Authorization": f"Bearer {api_key}",
                "Content-Type": "application/json",
            },
            json={"input": {"text": text}},
            timeout=120,
        )
        response.raise_for_status()
        data = response.json()

    except requests.RequestException:
        logger.exception("RunPod request failed")
        raise HTTPException(
            status_code=503,
            detail="AI classification service is temporarily unavailable.",
        )

    if data.get("status") != "COMPLETED":
        logger.warning("RunPod job not completed: %s", data.get("status"))
        raise HTTPException(
            status_code=503,
            detail="AI model is starting or the prediction is still queued. Please try again.",
        )

    output = data.get("output")

    if not isinstance(output, dict) or "predicted_class" not in output:
        logger.error("Invalid RunPod prediction response: %s", data.get("status"))
        raise HTTPException(
            status_code=502,
            detail="AI service returned an invalid prediction.",
        )

    category = str(output["predicted_class"])
    confidence = round(float(output["confidence"]) * 100, 2)

    top_categories = [{"category": category, "confidence": confidence}]

    return category, confidence, top_categories



logger = logging.getLogger("resume")

RESUME_RESULT_DIR = BASE_DIR / "data" / "resume_results"
RESUME_RESULT_DIR.mkdir(parents=True, exist_ok=True)

_STORED_NAME_RE = re.compile(r"^[0-9a-f]{32}\.(pdf|docx)$")


def _result_path(request: Request):
    """Result file for the signed-in user, or None if not signed in."""
    user = request.session.get("user")
    if not user or not user.get("sub"):
        return None
    key = hashlib.sha256(str(user["sub"]).encode()).hexdigest()
    return RESUME_RESULT_DIR / f"{key}.json"


def _delete_stored_file(stored_name):
    """Delete a saved resume, only if the name is exactly what we generated."""
    if stored_name and _STORED_NAME_RE.match(stored_name):
        (RESUME_STORAGE_DIR / stored_name).unlink(missing_ok=True)


def _looks_like_resume_file(extension: str, content: bytes) -> bool:
    if extension == ".pdf":
        return content.startswith(b"%PDF")
    if extension == ".docx":
        return content.startswith(b"PK\x03\x04")
    return False


def _run_analysis(filename: str, content: bytes) -> dict:
    """Blocking work (text extraction + model). Run in a threadpool."""
    text = extract_resume_text(filename, content)
    category, confidence, top_categories = predict_resume_category(text)

    words = re.findall(r"\b[\w+#.-]+\b", text)

    sections = {
        "contact": bool(re.search(r"\b[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}\b", text)),
        "education": bool(
            re.search(r"\b(education|degree|diploma|university|college)\b", text, re.I)
        ),
        "experience": bool(
            re.search(
                r"\b(experience|employment|work history|internship)\b", text, re.I
            )
        ),
        "skills": bool(
            re.search(r"\b(skills|technical skills|competencies)\b", text, re.I)
        ),
        "projects": bool(
            re.search(r"\b(projects|achievements|certifications)\b", text, re.I)
        ),
    }

    messages = {
        "contact": "Add a professional email address to your contact details.",
        "education": "Include your education and relevant qualifications.",
        "experience": "Add work experience, internships, or relevant experience.",
        "skills": "Add a clearly labelled skills section with relevant skills.",
        "projects": "Consider adding projects, achievements, or certifications.",
    }
    suggestions = [messages[k] for k, present in sections.items() if not present]

    return {
        "filename": filename,
        "predicted_category": category,
        "model_confidence": confidence,
        "top_categories": top_categories,
        "quality_score": round(sum(sections.values()) / len(sections) * 100),
        "score_type": "section_completeness_heuristic",
        "metrics": {k: 100 if v else 0 for k, v in sections.items()},
        "suggestions": suggestions,
        "word_count": len(words),
        "analysed_at": datetime.now(timezone.utc).isoformat(),
    }


def _quiz_path(request: Request):
    """Quiz result file for the signed-in user, or None if not signed in."""
    user = request.session.get("user")
    if not user or not user.get("sub"):
        return None
    key = hashlib.sha256(str(user["sub"]).encode()).hexdigest()
    return QUIZ_RESULT_DIR / f"{key}.json"


def _load_quiz_result(request: Request):
    path = _quiz_path(request)
    if not path or not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (ValueError, OSError):
        return None


def _save_quiz_result(request: Request, record: dict) -> bool:
    path = _quiz_path(request)
    if not path:
        return False
    try:
        tmp = path.with_suffix(".tmp")
        tmp.write_text(json.dumps(record), encoding="utf-8")
        tmp.replace(path)
        return True
    except OSError:
        logger.exception("Could not save quiz result")
        return False


PLAN_DIR = DATA_DIR / "plans"
PLAN_DIR.mkdir(parents=True, exist_ok=True)


def _plan_path(request: Request):
    user = request.session.get("user")
    if not user or not user.get("sub"):
        return None
    return PLAN_DIR / (hashlib.sha256(str(user["sub"]).encode()).hexdigest() + ".json")


def _load_plan(request: Request) -> list:
    p = _plan_path(request)
    try:
        return json.loads(p.read_text(encoding="utf-8")) if p and p.exists() else []
    except (ValueError, OSError):
        return []


def _resume_text_for(request: Request):
    """Text of the user's saved resume, or None if they haven't uploaded one."""
    path = _result_path(request)
    if not path or not path.exists():
        return None
    try:
        stored = json.loads(path.read_text(encoding="utf-8")).get("stored_file")
        if not stored or not _STORED_NAME_RE.match(stored):
            return None
        f = RESUME_STORAGE_DIR / stored
        return extract_resume_text(stored, f.read_bytes()) if f.exists() else None
    except Exception:
        return None


class PlanItem(BaseModel):
    code: str
    title: str
    missing_skills: list[str] = []


def _reskill_path(request: Request):
    user = request.session.get("user")
    if not user or not user.get("sub"):
        return None
    return RESKILL_DIR / (
        hashlib.sha256(str(user["sub"]).encode()).hexdigest() + ".json"
    )


def _load_reskill(path) -> dict:
    try:
        return (
            json.loads(path.read_text(encoding="utf-8"))
            if path and path.exists()
            else {}
        )
    except (ValueError, OSError):
        return {}


class ReskillRequest(BaseModel):
    code: str
    level: Literal["Beginner", "Intermediate", "Advanced"] = "Beginner"
    hours: int = Field(5, ge=1, le=15)
    style: str = Field("Mixed", max_length=60)
    free_only: bool = True
    target_date: str | None = Field(None, max_length=10)  
    force: bool = False 




@app.post("/api/resume/analyze")
async def analyze_resume(request: Request, file: UploadFile = File(...)):
    original_name = Path(file.filename or "").name
    extension = Path(original_name).suffix.lower()

    if extension not in ALLOWED_RESUME_EXTENSIONS:
        raise HTTPException(
            status_code=400, detail="Only PDF and DOCX files are supported."
        )

    content = await file.read(MAX_RESUME_SIZE + 1)

    if not content:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")
    if len(content) > MAX_RESUME_SIZE:
        raise HTTPException(
            status_code=413, detail="File exceeds the 5 MB upload limit."
        )
    if not _looks_like_resume_file(extension, content):
        raise HTTPException(
            status_code=400,
            detail="The file content does not match a valid PDF or DOCX.",
        )

    try:
        result = await run_in_threadpool(_run_analysis, original_name, content)
    except HTTPException:
        raise
    except Exception:
        logger.exception("Resume analysis failed")
        raise HTTPException(
            status_code=500, detail="Resume analysis failed. Please try again."
        )

    result_path = _result_path(request)
    result["saved"] = False

    if result_path:
        try:
            old_stored = None
            if result_path.exists():
                try:
                    old_stored = json.loads(
                        result_path.read_text(encoding="utf-8")
                    ).get("stored_file")
                except (ValueError, OSError):
                    old_stored = None

            stored_name = f"{uuid.uuid4().hex}{extension}"
            (RESUME_STORAGE_DIR / stored_name).write_bytes(content)

            result["saved"] = True
            result_path.write_text(
                json.dumps({"result": result, "stored_file": stored_name}),
                encoding="utf-8",
            )
            _delete_stored_file(old_stored)
        except OSError:
            logger.exception("Could not save resume result")
            result["saved"] = False

    return JSONResponse(result, headers={"Cache-Control": "no-store"})


@app.get("/api/resume/latest")
async def latest_resume_result(request: Request):
    path = _result_path(request)
    data = None
    if path and path.exists():
        try:
            data = json.loads(path.read_text(encoding="utf-8")).get("result")
        except (ValueError, OSError):
            data = None
    return JSONResponse({"result": data}, headers={"Cache-Control": "no-store"})


@app.delete("/api/resume/latest")
async def delete_resume_result(request: Request):
    path = _result_path(request)
    if path and path.exists():
        try:
            stored = json.loads(path.read_text(encoding="utf-8")).get("stored_file")
        except (ValueError, OSError):
            stored = None
        _delete_stored_file(stored)
        path.unlink(missing_ok=True)
    return JSONResponse({"ok": True}, headers={"Cache-Control": "no-store"})


app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")

templates = Jinja2Templates(directory=BASE_DIR / "templates")
try:
    from backend.chatbot_routes import router as chatbot_router
except ImportError: 
    from chatbot_routes import router as chatbot_router
app.include_router(chatbot_router)

app.add_middleware(
    SessionMiddleware,
    secret_key=os.environ["SESSION_SECRET"],
    same_site="lax",
    https_only=True,
)


oauth = OAuth()
oauth.register(
    name="google",
    client_id=os.getenv("GOOGLE_CLIENT_ID"),
    client_secret=os.getenv("GOOGLE_CLIENT_SECRET"),
    server_metadata_url="https://accounts.google.com/.well-known/openid-configuration",
    client_kwargs={"scope": "openid email profile"},
)


@app.get("/", response_class=HTMLResponse)
async def index_page(request: Request):
    return templates.TemplateResponse(
        request=request, name="index.html", context={"request": request}
    )


@app.get("/login")
async def login(request: Request):
    if not os.getenv("GOOGLE_CLIENT_ID") or not os.getenv("GOOGLE_CLIENT_SECRET"):
        return HTMLResponse(
            "<h1>Google sign-in needs configuration</h1>"
            "<p>Add GOOGLE_CLIENT_ID and GOOGLE_CLIENT_SECRET in Replit Secrets, "
            "then register this site's /auth URL as a Google OAuth redirect URI.</p>"
            '<p><a href="/">Back to welcome page</a></p>',
            status_code=503,
        )
    redirect_uri = request.url_for("auth")

    return await oauth.google.authorize_redirect(request, redirect_uri)


@app.get("/auth")
async def auth(request: Request):
    if not os.getenv("GOOGLE_CLIENT_ID") or not os.getenv("GOOGLE_CLIENT_SECRET"):
        return RedirectResponse(url="/login", status_code=303)
    token = await oauth.google.authorize_access_token(request)

    user = token.get("userinfo")

    if not user:
        return RedirectResponse(url="/login", status_code=303)

    request.session["user"] = {
        "sub": user.get("sub"),
        "name": user.get("name"),
        "email": user.get("email"),
        "picture": user.get("picture"),
    }

    return RedirectResponse(url="/home", status_code=303)


@app.get("/home", response_class=HTMLResponse)
async def homepage(request: Request):
    user = request.session.get("user")

    if not user:
        return RedirectResponse(url="/login", status_code=303)

    return templates.TemplateResponse(
        request=request,
        name="homepage.html",
        context={"request": request, "user": user},
    )


@app.get("/about", response_class=HTMLResponse)
async def about_us(request: Request):
    return templates.TemplateResponse(
        request=request, name="about_us.html", context={"request": request}
    )


@app.get("/mabout", response_class=HTMLResponse)
async def mabout_us(request: Request):
    return templates.TemplateResponse(
        request=request, name="mabout_us.html", context={"request": request}
    )


@app.get("/profile", response_class=HTMLResponse)
async def profile(request: Request):
    return templates.TemplateResponse(
        request=request, name="profile.html", context={"request": request}
    )


@app.get("/logout")
async def logout(request: Request):
    request.session.clear()

    return RedirectResponse(url="/", status_code=303)


@app.get("/career-quiz", response_class=HTMLResponse)
async def career_quiz(request: Request):
    data = {
        "questions": [
            {"n": n, "text": text, "cat": QUESTION_CATEGORY[n]}
            for n, text in QUESTIONS.items()
        ],
        "cats": [{"name": c, "letter": RIASEC_LETTERS[c]} for c in RIASEC_COLUMNS],
        "saved": _load_quiz_result(request), 
    }
    return templates.TemplateResponse(
        request=request,
        name="career_matches.html",
        context={"request": request, "data": data},
    )


@app.post("/api/quiz/submit")
def quiz_submit(payload: QuizSubmission, request: Request):
    try:
        answers = {q: ("Y" if v >= 4 else "N") for q, v in payload.answers.items()}
        scores, code = score_riasec(answers)
        results = match_careers(scores, top_n=10)
    except ValueError as e:
        return JSONResponse({"error": str(e)}, status_code=400)
    except FileNotFoundError as e:
        return JSONResponse({"error": f"Dataset missing: {e}"}, status_code=500)

    careers = json.loads(results.to_json(orient="records"))
    careers = apply_ai_salaries(careers)

    record = {
        "scores": scores,
        "code": code,
        "careers": careers,
        "answers": {str(q): v for q, v in payload.answers.items()},
        "completed_at": datetime.now(timezone.utc).isoformat(),
    }
    saved = _save_quiz_result(request, record)

    return JSONResponse(
        {**record, "saved": saved}, headers={"Cache-Control": "no-store"}
    )


@app.get("/coach", response_class=HTMLResponse)
async def coach(request: Request):
    return templates.TemplateResponse(
        request=request, name="resume_coach.html", context={"request": request}
    )


@app.get("/futuresk", response_class=HTMLResponse)
async def futuresk(request: Request):
    return templates.TemplateResponse(
        request=request, name="future_skills.html", context={"request": request}
    )


@app.get("/careerde", response_class=HTMLResponse)
async def careerde(request: Request):
    return templates.TemplateResponse(
        request=request, name="career_detail.html", context={"request": request}
    )


@app.get("/api/career/{code}")
def career_api(code: str, request: Request):
    quiz = _load_quiz_result(request)
    if not quiz:
        return JSONResponse(
            {"error": "Take the Career Quiz (while signed in) to see career details."},
            status_code=404,
        )
    if code == "top":
        code = quiz["careers"][0]["O*NET-SOC Code"]
    try:
        data = build_detail(code, quiz["scores"], _resume_text_for(request))
    except KeyError:
        return JSONResponse({"error": "Career not found."}, status_code=404)
    except Exception as e:
        logger.exception("Career detail failed")
        return JSONResponse(
            {"error": f"Server error: {type(e).__name__}: {e}"}, status_code=500
        )
    data["saved"] = any(p["code"] == code for p in _load_plan(request))
    return JSONResponse(data, headers={"Cache-Control": "no-store"})


@app.get("/api/plan")
def plan_get(request: Request):
    return JSONResponse(
        {"plan": _load_plan(request)}, headers={"Cache-Control": "no-store"}
    )


@app.post("/api/plan")
def plan_save(item: PlanItem, request: Request):
    path = _plan_path(request)
    if not path:
        return JSONResponse({"error": "Sign in to save a plan."}, status_code=401)
    plan = [p for p in _load_plan(request) if p["code"] != item.code]
    plan.append(
        {**item.model_dump(), "saved_at": datetime.now(timezone.utc).isoformat()}
    )
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(plan), encoding="utf-8")
    tmp.replace(path)
    return JSONResponse({"ok": True})


@app.get("/reskilling", response_class=HTMLResponse)
async def reskilling(request: Request):
    return templates.TemplateResponse(
        request=request, name="reskilling.html", context={"request": request}
    )


@app.get("/api/future-skills/wage-trends")
def get_wage_trends():
    file_path = DATA_DIR / "wages_occupation.csv"

    if not file_path.exists():
        raise HTTPException(
            status_code=404, detail="Occupation wage dataset not found."
        )

    try:
        df = pd.read_csv(file_path)
    except Exception:
        logger.exception("Could not read wage dataset")
        raise HTTPException(status_code=500, detail="Could not read the wage dataset.")

    df.columns = [str(c).strip().lower() for c in df.columns]
    df["year"] = pd.to_datetime(df["date"], errors="coerce").dt.year

    df = df[(df["sex"] == "both") & (df["occupation"] != "overall")]
    df = df.dropna(subset=["year", "wages_median"])
    df["year"] = df["year"].astype(int)

    latest_year = int(df["year"].max())
    start_year = latest_year - 3

    occupations = []
    for name, grp in df.groupby("occupation"):
        by_year = dict(zip(grp["year"], grp["wages_median"]))
        if latest_year not in by_year or start_year not in by_year:
            continue
        latest = float(by_year[latest_year])
        start = float(by_year[start_year])
        occupations.append(
            {
                "name": name,
                "latest": latest,
                "change_pct": round((latest - start) / start * 100, 1),
                "series": [
                    {"year": int(y), "median": float(v)}
                    for y, v in sorted(by_year.items())
                ],
            }
        )

    occupations.sort(key=lambda o: o["change_pct"], reverse=True)

    return JSONResponse(
        {
            "latest_year": latest_year,
            "start_year": start_year,
            "occupations": occupations,
        },
        headers={"Cache-Control": "no-store"},
    )


@app.get("/api/reskilling/plan")
def reskilling_plan_get(code: str, request: Request):
    path = _reskill_path(request)
    if not path:
        return JSONResponse({"error": "Sign in to see your plan."}, status_code=401)
    entry = _load_reskill(path).get(code)
    return JSONResponse(
        {
            "plan": entry["plan"] if entry else None,
            "prefs": entry["prefs"] if entry else None,
        },
        headers={"Cache-Control": "no-store"},
    )

@app.post("/api/reskilling/plan")
def reskilling_plan_make(req: ReskillRequest, request: Request):
    path = _reskill_path(request)
    if not path:
        return JSONResponse({"error": "Sign in to create a plan."}, status_code=401)

    item = next((p for p in _load_plan(request) if p["code"] == req.code), None)
    if not item:
        return JSONResponse(
            {"error": "Save this career from Career Matches first."}, status_code=404
        )
    skills = item.get("missing_skills") or []
    if not skills:
        return JSONResponse(
            {
                "error": "No skill gaps were saved for this career, so there is nothing to learn."
            },
            status_code=400,
        )

    prefs = req.model_dump(exclude={"force", "code"})
    key = hashlib.sha256(
        json.dumps([prefs, skills], sort_keys=True).encode()
    ).hexdigest()

    store = _load_reskill(path)
    cached = store.get(req.code)
    if cached and cached.get("key") == key and not req.force:
        return JSONResponse(
            {"plan": cached["plan"], "prefs": cached["prefs"], "cached": True},
            headers={"Cache-Control": "no-store"},
        )

    try:
        plan = generate_plan(
            career=item["title"],
            skills=skills,
            level=req.level,
            hours_per_week=req.hours,
            style=req.style,
            free_only=req.free_only,
            target_date=req.target_date,
        )
    except ValueError as e:
        return JSONResponse({"error": str(e)}, status_code=400)

    store[req.code] = {
        "key": key,
        "prefs": prefs,
        "plan": plan,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(store), encoding="utf-8")
    tmp.replace(path)

    return JSONResponse(
        {"plan": plan, "prefs": prefs, "cached": False},
        headers={"Cache-Control": "no-store"},
    )




class LessonRef(BaseModel):
    code: str
    skill: str = Field(..., max_length=80)
    refresh: bool = False


class QuizResult(BaseModel):
    code: str
    skill: str = Field(..., max_length=80)
    score: int = Field(..., ge=0, le=20)
    total: int = Field(..., ge=1, le=20)


def _lesson_ctx(request: Request, code: str, skill: str):
    """Returns (career_title, prefs, error_response). Skill must be one the user saved."""
    uid = _uid(request)
    if not uid:
        return None, None, JSONResponse({"error": "Sign in first."}, status_code=401)
    item = next((p for p in _load_plan(request) if p["code"] == code), None)
    if not item:
        return (
            None,
            None,
            JSONResponse({"error": "Save this career first."}, status_code=404),
        )
    entry = _load_reskill(_reskill_path(request)).get(code, {})
    plan_skills = {
        m.get("skill")
        for ph in (entry.get("plan") or {}).get("phases", [])
        for m in ph.get("modules", [])
    }
    allowed = set(item.get("missing_skills") or []) | plan_skills
    if skill not in allowed:
        return (
            None,
            None,
            JSONResponse({"error": "That skill is not in your plan."}, status_code=400),
        )
    saved = entry.get("prefs") or {}
    prefs = {
        "level": saved.get("level", "Beginner"),
        "style": saved.get("style", "Mixed"),
        "hours": saved.get("hours", 5),
        "free_only": saved.get("free_only", True),
    }
    return item["title"], prefs, None


def _lesson_cache(uid, kind, code, skill, prefs):
    k = hashlib.sha256(
        json.dumps([kind, code, skill, prefs], sort_keys=True).encode()
    ).hexdigest()[:20]
    return LESSON_DIR / f"{uid}_{kind}_{k}.json"


@app.get("/api/reskilling/lesson")
def lesson_get(code: str, skill: str, request: Request, refresh: bool = False):
    career, prefs, err = _lesson_ctx(request, code, skill)
    if err:
        return err
    cache = _lesson_cache(_uid(request), "lesson", code, skill, prefs)
    if cache.exists() and not refresh:
        return JSONResponse(
            json.loads(cache.read_text(encoding="utf-8")),
            headers={"Cache-Control": "no-store"},
        )
    try:
        lesson = generate_lesson(
            career,
            skill,
            prefs["level"],
            prefs["style"],
            prefs["hours"],
            prefs["free_only"],
        )
    except ValueError as e:
        return JSONResponse({"error": str(e)}, status_code=502)
    cache.write_text(json.dumps(lesson), encoding="utf-8")
    return JSONResponse(lesson, headers={"Cache-Control": "no-store"})


@app.post("/api/reskilling/quiz")
def lesson_quiz(ref: LessonRef, request: Request):
    career, prefs, err = _lesson_ctx(request, ref.code, ref.skill)
    if err:
        return err
    cache = _lesson_cache(_uid(request), "quiz", ref.code, ref.skill, prefs)
    if cache.exists() and not ref.refresh:
        return JSONResponse(
            {"questions": json.loads(cache.read_text(encoding="utf-8"))},
            headers={"Cache-Control": "no-store"},
        )
    summary = ""
    lc = _lesson_cache(_uid(request), "lesson", ref.code, ref.skill, prefs)
    if lc.exists():
        summary = json.loads(lc.read_text(encoding="utf-8")).get("summary", "")
    try:
        qs = generate_quiz(career, ref.skill, prefs["level"], summary)
    except ValueError as e:
        return JSONResponse({"error": str(e)}, status_code=502)
    cache.write_text(json.dumps(qs), encoding="utf-8")
    return JSONResponse({"questions": qs}, headers={"Cache-Control": "no-store"})


def _progress_file(request: Request):
    uid = _uid(request)
    return LESSON_DIR / f"{uid}_progress.json" if uid else None


@app.post("/api/reskilling/quiz/result")
def quiz_result_save(r: QuizResult, request: Request):
    path = _progress_file(request)
    if not path:
        return JSONResponse({"error": "Sign in first."}, status_code=401)
    try:
        data = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    except (OSError, ValueError):
        data = {}
    passed = r.score / r.total >= 0.6
    prev = data.setdefault(r.code, {}).get(r.skill, {})
    data[r.code][r.skill] = {
        "score": r.score,
        "total": r.total,
        "passed": passed or bool(prev.get("passed")), 
        "best": max(r.score / r.total, prev.get("best", 0)),
        "at": datetime.now(timezone.utc).isoformat(),
    }
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(data), encoding="utf-8")
    tmp.replace(path)
    return {"ok": True, "passed": passed}


@app.get("/api/reskilling/progress")
def lesson_progress(request: Request):
    path = _progress_file(request)
    if not path:
        return JSONResponse({"error": "Sign in first."}, status_code=401)
    try:
        data = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    except (OSError, ValueError):
        data = {}
    return JSONResponse({"progress": data}, headers={"Cache-Control": "no-store"})



@app.get("/sdg", response_class=HTMLResponse)
async def sdg_page(request: Request):
    return templates.TemplateResponse(
        request=request, name="sdg_impact.html", context={"request": request}
    )


def _career_context(request: Request, code: str):
    """Pick the career: ?code=... > latest saved plan item > top quiz match."""
    quiz = _load_quiz_result(request)
    if not quiz:
        return None
    plan = _load_plan(request)
    item = None
    if code == "auto":
        item = plan[-1] if plan else None 
        code = item["code"] if item else quiz["careers"][0]["O*NET-SOC Code"]
    else:
        item = next((p for p in plan if p["code"] == code), None)

    if item:
        title, missing = item["title"], item.get("missing_skills") or []
    else:
        d = build_detail(code, quiz["scores"], _resume_text_for(request))
        title = d.get("title") or d.get("name") or d.get("career") or code
        missing = d.get("missing_skills") or d.get("skill_gaps") or []
    have = []  # TODO: fill from build_detail's "skills you have" key
    return code, title, have, missing


@app.get("/api/sdg")
def sdg_api(request: Request, code: str = "auto", refresh: bool = False):
    if not request.session.get("user"):
        return JSONResponse(
            {"error": "Sign in to see your SDG impact."}, status_code=401
        )
    try:
        ctx = _career_context(request, code)
    except KeyError:
        return JSONResponse({"error": "Career not found."}, status_code=404)
    if not ctx:
        return JSONResponse(
            {"error": "Take the Career Quiz first so we know your career."},
            status_code=404,
        )
    code, title, have, missing = ctx

    user_key = hashlib.sha256(str(request.session["user"]["sub"]).encode()).hexdigest()
    key = hashlib.sha256(
        json.dumps([code, have, missing], sort_keys=True).encode()
    ).hexdigest()[:16]
    cache = SDG_DIR / f"{user_key}_{key}.json"
    if cache.exists() and not refresh:
        return JSONResponse(
            json.loads(cache.read_text(encoding="utf-8")),
            headers={"Cache-Control": "no-store"},
        )

    result = sdg_analyse(title, have, missing)
    result["career"] = {"code": code, "title": title}
    result["generated_at"] = datetime.now(timezone.utc).isoformat()
    cache.write_text(json.dumps(result), encoding="utf-8")
    return JSONResponse(result, headers={"Cache-Control": "no-store"})


def _deepseek_json(prompt: str, timeout: int = 45):
    key = os.getenv("DEEPSEEK_API_KEY")
    if not key:
        return None
    try:
        r = requests.post(
            "https://api.deepseek.com/chat/completions",
            headers={"Authorization": f"Bearer {key}"},
            json={
                "model": DEEPSEEK_MODEL,
                "temperature": 0.2,
                "response_format": {"type": "json_object"},
                "messages": [{"role": "user", "content": prompt}],
            },
            timeout=timeout,
        )
        r.raise_for_status()
        return json.loads(r.json()["choices"][0]["message"]["content"])
    except Exception:
        logger.exception("DeepSeek call failed")
        return None


def _sc_series():
    p = DATA_DIR / "wages_occupation.csv"
    if not p.exists():
        raise HTTPException(
            status_code=404, detail="Occupation wage dataset not found."
        )
    df = pd.read_csv(p)
    df.columns = [str(c).strip().lower() for c in df.columns]
    df["year"] = pd.to_datetime(df["date"], errors="coerce").dt.year
    df = df[(df["sex"] == "both") & (df["occupation"] != "overall")]
    df = df.dropna(subset=["year", "wages_median"])
    df["year"] = df["year"].astype(int)
    latest_year = int(df["year"].max())
    series = {}
    for name, g in df.groupby("occupation"):
        by = dict(sorted(zip(g["year"], g["wages_median"].astype(float))))
        if latest_year in by and len(by) >= 5:
            series[name] = by
    return latest_year, series


def _fit(ys, vs, future):
    """Log-linear trend fitted to the data, with a real 90% prediction interval."""
    x = np.array(ys, float)
    y = np.log(np.array(vs, float))
    n = len(x)
    b, a = np.polyfit(x, y, 1)
    s = float(np.sqrt(((y - (a + b * x)) ** 2).sum() / max(n - 2, 1)))
    sxx = ((x - x.mean()) ** 2).sum()
    out = []
    for f in future:
        se = s * math.sqrt(1 + 1 / n + (f - x.mean()) ** 2 / sxx)
        out.append((a + b * f, se))
    return float(b), out


def _backtest(by):
    """Hide the last 2 years, forecast them, report average % error."""
    ys = list(by)
    if len(ys) < 6:
        return None
    train = ys[:-2][-8:]
    _, pred = _fit(train, [by[y] for y in train], ys[-2:])
    errs = [abs(math.exp(m) - by[y]) / by[y] for (m, _), y in zip(pred, ys[-2:])]
    return round(float(np.mean(errs)) * 100, 1)


def _exposure(names):
    """AI exposure per occupation, estimated by DeepSeek once and cached on disk."""
    try:
        cache = json.loads(SC_EXPOSURE_FILE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        cache = {}
    missing = [n for n in names if n not in cache]
    if missing:
        d = _deepseek_json(
            "You are a labour economist. For each Malaysian occupation group below, give rough 0-100 "
            "judgements for the next 5 years: 'automation' = share of its tasks AI could replace, "
            "'augmentation' = share where AI could raise worker productivity. Base this on widely known "
            "research on generative-AI task exposure. Do not quote statistics. 'why' = one plain sentence. "
            'Return JSON: {"items":[{"name":"<exact name>","automation":n,"augmentation":n,"why":"..."}]}. '
            "Occupation groups: " + json.dumps(missing)
        )
        for it in (d or {}).get("items", []):
            try:
                if it["name"] in missing:
                    cache[it["name"]] = {
                        "automation": int(max(0, min(100, float(it["automation"])))),
                        "augmentation": int(
                            max(0, min(100, float(it["augmentation"])))
                        ),
                        "why": str(it.get("why", ""))[:220],
                        "by": f"DeepSeek ({DEEPSEEK_MODEL})",
                        "date": date.today().isoformat(),
                    }
            except (KeyError, ValueError, TypeError):
                continue
        SC_EXPOSURE_FILE.write_text(json.dumps(cache), encoding="utf-8")
    return {n: cache[n] for n in names if n in cache}


def _sc_build() -> dict:
    latest_year, series = _sc_series()
    fits = {}
    for name, by in series.items():
        ys = list(by)[-8:]
        b, _ = _fit(ys, [by[y] for y in ys], [])
        fits[name] = (ys, b)
    spread = max(
        float(np.std([b for _, b in fits.values()])), 0.005
    ) 
    expo = _exposure(list(series))

    out = []
    for name, by in series.items():
        ys, b = fits[name]
        _, base = _fit(
            ys, [by[y] for y in ys], [latest_year + t for t in range(1, SC_YEARS + 1)]
        )
        e = expo.get(name)
        net = (e["augmentation"] - e["automation"]) / 100 if e else 0.0
        latest = by[latest_year]
        scen = {}
        for key, speed in SC_SPEED.items():
            shift = net * speed * spread 
            pts = []
            for t, (m, se) in enumerate(base, 1):
                m2 = m + shift * t
                pts.append(
                    {
                        "year": latest_year + t,
                        "mid": round(math.exp(m2), 1),
                        "lo": round(math.exp(m2 - 1.645 * se), 1),
                        "hi": round(math.exp(m2 + 1.645 * se), 1),
                    }
                )
            scen[key] = {
                "change_pct": round((pts[-1]["mid"] / latest - 1) * 100, 1),
                "points": pts,
            }
        out.append(
            {
                "name": name,
                "latest": latest,
                "hist_growth_pct": round((math.exp(b) - 1) * 100, 2),
                "ai_adjustment_pp": round(net * SC_SPEED["moderate"] * spread * 100, 2),
                "exposure": e,
                "backtest_mape_pct": _backtest(by),
                "history": [{"year": int(y), "median": v} for y, v in by.items()],
                "scenarios": scen,
            }
        )
    return {
        "latest_year": latest_year,
        "source": "DOSM occupational wages (data/wages_occupation.csv) + AI exposure estimated by DeepSeek",
        "method": "Trend = log-linear fit to the last 8 years of median wages (a proxy for demand); the shaded "
        "band is the 90% prediction interval from how well that fit matched the data. Each scenario "
        "adds an AI effect = (augmentation - automation) x adoption speed x the real spread of "
        "growth between occupations. Exposure scores are an LLM judgement, not an official statistic.",
        "occupations": out,
    }


def _sc_user_career(request: Request):
    """The signed-in user's target career: latest saved plan item, else top quiz match."""
    quiz = _load_quiz_result(request)
    if not quiz:
        return None, []
    plan = _load_plan(request)
    if plan:
        return plan[-1]["title"], plan[-1].get("missing_skills") or []
    c = quiz["careers"][0]
    return next((str(v) for k, v in c.items() if "title" in k.lower()), None), []


@app.get("/scenarios", response_class=HTMLResponse)
async def scenarios_page(request: Request):
    return templates.TemplateResponse(
        request=request, name="scenarios.html", context={"request": request}
    )


@app.get("/api/scenarios")
def scenarios_api():
    return JSONResponse(_sc_build(), headers={"Cache-Control": "no-store"})


@app.get("/api/scenarios/me")
def scenarios_me(request: Request):
    title, skills = _sc_user_career(request)
    if not title:
        return {"occupation": None}
    names = [o["name"] for o in _sc_build()["occupations"]]
    d = _deepseek_json(
        f'Pick the ONE occupation group from the list that best fits the career "{title}". '
        'Return JSON {"occupation": "<exact item from the list>"}. List: '
        + json.dumps(names)
    )
    occ = (d or {}).get("occupation")
    return {
        "career": title,
        "occupation": occ if occ in names else None,
        "skills": skills,
    }


@app.get("/api/scenarios/insight")
def scenarios_insight(request: Request, occupation: str, personal: bool = False):
    data = _sc_build()
    occ = next((o for o in data["occupations"] if o["name"] == occupation), None)
    if not occ:
        return JSONResponse({"error": "Occupation not found."}, status_code=404)
    m, f = occ["scenarios"]["moderate"], occ["scenarios"]["fast"]
    facts = {
        "occupation": occupation,
        "latest_year": data["latest_year"],
        "years": SC_YEARS,
        "moderate_change_pct": m["change_pct"],
        "fast_change_pct": f["change_pct"],
        "historical_growth_pct_per_year": occ["hist_growth_pct"],
        "ai_exposure": occ["exposure"],
        "backtest_error_pct": occ["backtest_mape_pct"],
    }
    if personal:
        title, skills = _sc_user_career(request)
        facts["user_target_career"], facts["user_skills_being_built"] = title, skills

    fallback = [
        f"Median pay for this group is estimated to change {m['change_pct']}% (moderate AI) or {f['change_pct']}% (fast AI) over {SC_YEARS} years.",
        f"Past wage growth was about {occ['hist_growth_pct']}% a year; the shaded band shows how uncertain the forecast is.",
        "These are estimates from past data and an AI judgement, not guarantees. Compare futures, then talk to a career adviser.",
    ]
    d = _deepseek_json(
        "Write exactly 3 short plain-English takeaways (max 28 words each) for a Malaysian worker reading this "
        "scenario forecast. If user_target_career is present, make the third one about how their skills relate "
        "to the scenarios. Use ONLY the facts given: do not invent jobs, courses, employers or statistics. "
        'Say these are estimates. Return JSON {"takeaways":["...","...","..."]}.\n'
        + json.dumps(facts)
    )
    items = [str(x) for x in (d or {}).get("takeaways", [])][:3]
    if len(items) == 3:
        return {"takeaways": items, "by": f"DeepSeek ({DEEPSEEK_MODEL})"}
    return {"takeaways": fallback, "by": "template"}



def _uid(request: Request):
    u = request.session.get("user")
    return (
        hashlib.sha256(str(u["sub"]).encode()).hexdigest()
        if u and u.get("sub")
        else None
    )


class Prefs(BaseModel):
    display_name: str = Field("", max_length=60)
    language: Literal["en", "ms", "zh"] = "en"
    weekly_hours: int = Field(5, ge=1, le=15)


def _load_prefs(request: Request) -> dict:
    uid = _uid(request)
    base = Prefs().model_dump()
    try:
        return {
            **base,
            **json.loads((SETTINGS_DIR / f"{uid}.json").read_text(encoding="utf-8")),
        }
    except (OSError, ValueError, TypeError):
        return base


def _apply_saved_name(request: Request):
    """Call this inside /auth right after request.session["user"] = {...}"""
    name = _load_prefs(request).get("display_name")
    if name:
        request.session["user"]["name"] = name


def _progress(request: Request) -> list:
    uid = _uid(request)
    rp, qp = _result_path(request), _quiz_path(request)
    return [
        {
            "key": "resume",
            "label": "Upload your CV",
            "href": "/profile",
            "done": bool(rp and rp.exists()),
        },
        {
            "key": "quiz",
            "label": "Take the career quiz",
            "href": "/career-quiz",
            "done": bool(qp and qp.exists()),
        },
        {
            "key": "plan",
            "label": "Save a career to your plan",
            "href": "/careerde",
            "done": bool(_load_plan(request)),
        },
        {
            "key": "reskilling",
            "label": "Build a reskilling plan",
            "href": "/reskilling",
            "done": bool(_load_reskill(_reskill_path(request))),
        },
        {
            "key": "sdg",
            "label": "See your SDG impact",
            "href": "/sdg",
            "done": any(SDG_DIR.glob(f"{uid}_*.json")),
        },
    ]


@app.get("/settings", response_class=HTMLResponse)
async def settings_page(request: Request):
    if not request.session.get("user"):
        return RedirectResponse(url="/login", status_code=303)
    return templates.TemplateResponse(
        request=request, name="settings.html", context={"request": request}
    )


@app.get("/api/settings")
def settings_get(request: Request):
    if not _uid(request):
        return JSONResponse({"error": "Sign in first."}, status_code=401)
    u = request.session["user"]
    steps = _progress(request)
    return JSONResponse(
        {
            "user": {
                "name": u.get("name"),
                "email": u.get("email"),
                "picture": u.get("picture"),
            },
            "prefs": _load_prefs(request),
            "steps": steps,
            "percent": round(sum(s["done"] for s in steps) / len(steps) * 100),
        },
        headers={"Cache-Control": "no-store"},
    )


@app.post("/api/settings")
def settings_save(p: Prefs, request: Request):
    uid = _uid(request)
    if not uid:
        return JSONResponse({"error": "Sign in first."}, status_code=401)
    data = p.model_dump()
    data["display_name"] = data["display_name"].strip()
    tmp = SETTINGS_DIR / f"{uid}.tmp"
    tmp.write_text(json.dumps(data), encoding="utf-8")
    tmp.replace(SETTINGS_DIR / f"{uid}.json")
    if data["display_name"]:
        request.session["user"]["name"] = data["display_name"]
    return {"ok": True, "prefs": data}


@app.delete("/api/settings/data")
def settings_delete(
    what: Literal["resume", "quiz", "plan", "reskilling", "sdg", "all"],
    request: Request,
):
    uid = _uid(request)
    if not uid:
        return JSONResponse({"error": "Sign in first."}, status_code=401)
    if what in ("resume", "all"):
        rp = _result_path(request)
        if rp and rp.exists():
            try:
                _delete_stored_file(
                    json.loads(rp.read_text(encoding="utf-8")).get("stored_file")
                )
            except (OSError, ValueError):
                pass
            rp.unlink(missing_ok=True)
    if what in ("quiz", "all"):
        _quiz_path(request).unlink(missing_ok=True)
    if what in ("plan", "all"):
        _plan_path(request).unlink(missing_ok=True)
    if what in ("reskilling", "all"):
        _reskill_path(request).unlink(missing_ok=True)
    if what in ("sdg", "all"):
        for f in SDG_DIR.glob(f"{uid}_*.json"):
            f.unlink(missing_ok=True)
    if what == "all":
        (SETTINGS_DIR / f"{uid}.json").unlink(missing_ok=True)
    if what == "all":
        from models import chat_store

        chat_store.delete_all(uid)
    return {"ok": True}


@app.get("/api/settings/export")
def settings_export(request: Request):
    if not _uid(request):
        return JSONResponse({"error": "Sign in first."}, status_code=401)
    rp = _result_path(request)
    try:
        resume = (
            json.loads(rp.read_text(encoding="utf-8")).get("result")
            if rp and rp.exists()
            else None
        )
    except (OSError, ValueError):
        resume = None
    dump = {
        "exported_at": datetime.now(timezone.utc).isoformat(),
        "profile": {k: request.session["user"].get(k) for k in ("name", "email")},
        "preferences": _load_prefs(request),
        "resume_analysis": resume,
        "quiz": _load_quiz_result(request),
        "career_plan": _load_plan(request),
        "reskilling": _load_reskill(_reskill_path(request)),
    }
    return Response(
        json.dumps(dump, indent=2),
        media_type="application/json",
        headers={"Content-Disposition": 'attachment; filename="kz-my-data.json"'},
    )
